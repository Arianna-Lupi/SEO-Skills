#!/usr/bin/env python3
"""Genera un brief de contenido por fila del mapa (content-map.json), en Markdown listo para Google Docs.

Secciones según la plantilla de brief de la comunidad y las Semanas 8 y 10 del diplomado
(ver references/plantillas-comunidad.md):
  blog          -> "Brief SEO de Blog"
  landing       -> "Brief de páginas de venta o Landing Page" (Home y servicios)
  ecommerce     -> "Brief de e-commerce" (productos y categorías; Semana 8 lección 10)
  actualizacion -> "Brief SEO de contenido existente" (Semana 10; acción = Actualizar)

El script NO redacta contenido creativo: rellena lo que existe en el mapa (keyword, URL, H1, metas,
secundarias, enlaces) y deja guías con el criterio del diplomado donde falta (metatítulo 50-60 c con la
keyword al inicio, metadescripción 120-155 c con CTA, H1 único, sin saltos H2 a H4). Lo que falta
dice `pendiente de dato`. Si pasas metas, cuenta los caracteres y avisa.

Salida: <salida>/briefs/{blog,landings,ecommerce,actualizacion}/<slug>.md, <salida>/briefs/00-indice.md,
y actualiza el campo `brief` del content-map.json. Resumen JSON por stdout.
.docx opcional con --docx (requiere python-docx; si falta, avisa y sigue con Markdown).

Uso:
  python3 generar_briefs.py --sitio ejemplo.com
  python3 generar_briefs.py --datos ./data --salida ./entregables --docx
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _plan_comun import PEND, asegurar_dir, ice, leer_json, slug  # noqa: E402

CARPETA = {"blog": "blog", "landing": "landings", "ecommerce": "ecommerce", "actualizacion": "actualizacion"}
GUIA_IMG = "Todas las imágenes deben tener nombre de archivo y texto ALT con la palabra clave principal o una variante natural."
GUIA_SEO = ("1. Contenido legible, párrafos de 2 o 3 líneas. 2. Un solo H1 con la keyword y sin saltar niveles (H2 a H4 está prohibido). "
            "3. Revisa la jerarquía a mano si usas IA: suele romperla. 4. Enlaces internos definidos aquí, no se los dejes al redactor.")


def tipo_brief(p):
    if p.get("accion") == "Actualizar":
        return "actualizacion"
    t = p.get("tipo")
    if t == "Blog":
        return "blog"
    if t in ("Ecommerce", "Categoría"):
        return "ecommerce"
    return "landing"


def v(p, k, defecto=PEND):
    x = p.get(k)
    return defecto if x in (None, "", []) else x


def fila(a, b):
    return "| %s | %s |" % (a, str(b).replace("\n", " ").replace("|", "/"))


def cuenta(texto, lo, hi):
    if texto in (None, "", PEND):
        return PEND
    n = len(texto)
    return "%s (%d caracteres%s)" % (texto, n, "" if lo <= n <= hi else ", FUERA de rango %d-%d" % (lo, hi))


def bloque_metas(p, actual=False):
    L = []
    if actual:
        L += [fila("Meta título actual", v(p, "meta_title_actual")), fila("Meta descripción actual", v(p, "meta_desc_actual"))]
    L += [fila("Meta título sugerido (50-60 c, keyword al inicio)", cuenta(p.get("meta_title"), 50, 60) if p.get("meta_title") else PEND + ": redáctalo con la keyword al inicio"),
          fila("Meta descripción sugerida (120-155 c, con CTA)", cuenta(p.get("meta_desc"), 120, 155) if p.get("meta_desc") else PEND + ": redáctala persuasiva y con CTA")]
    return L


def bloque_links(p):
    ent = p.get("enlaces_entrantes") or []
    sal = p.get("enlaces_salientes") or []
    L = ["", "### III. Enlaces internos", "", "**Páginas que enlazarán a ESTA página**", "", "| URL | Anchor |", "|---|---|"]
    L += ["| %s | %s |" % (e.get("url", PEND), e.get("anchor", PEND)) for e in ent] or ["| %s | %s |" % (PEND, PEND)]
    L += ["", "**Páginas a las que enlazará ESTA página** (usa el cluster: pilar y hermanas)", "", "| URL | Anchor |", "|---|---|"]
    L += ["| %s | %s |" % (e.get("url", PEND), e.get("anchor", PEND)) for e in sal] or ["| %s | %s |" % (PEND, PEND)]
    return L


def bloque_estructura(p):
    est = p.get("estructura") or []
    L = ["", "### Sugerencias de la estructura", ""]
    L += ["- " + e for e in est] or ["- H1: %s" % v(p, "h1", PEND + " (un solo H1 con la keyword)"),
                                     "  - H2: %s" % PEND, "    - H3: %s" % PEND, "  - H2: %s (responde una pregunta de People Also Ask)" % PEND]
    return L


def bloque_nlp(p):
    return ["", "### NLP Report", "", "| Elemento | Sugerencias |", "|---|---|", fila("Puntuación", PEND), fila("Contenido del reporte", PEND),
            fila("Conteo de palabras", v(p, "palabras")), fila("Ejemplos / inspiración (top de la SERP)", ", ".join(p.get("inspiracion") or []) or PEND)]


def brief(p, tipo, sitio):
    kw = v(p, "keyword")
    titulos = {"blog": "Brief SEO de Blog", "landing": "Brief de páginas de venta o Landing Page",
               "ecommerce": "Brief de contenido para e-commerce", "actualizacion": "Brief SEO de contenido existente"}
    L = ["# %s" % titulos[tipo], "", "Equipo de SEO · %s" % sitio, "", "| | |", "|---|---|", fila("Fecha", dt.date.today().strftime("%d/%m/%Y")),
         fila("Brief", kw), fila("Escritor", v(p, "responsable")), fila("Fecha de entrega", v(p, "fecha")), fila("Cluster", "%s · %s" % (v(p, "cluster"), v(p, "rol"))),
         fila("Puntuación ICE", ice(p) if ice(p) is not None else PEND), "", "> Regla de datos: ningún volumen, KD ni cifra se inventa. Lo que falta dice `pendiente de dato`.", ""]
    if tipo == "blog":
        L += ["### I. Elementos de la página", "", "| Elemento | Sugerencias |", "|---|---|",
              fila("Intención de búsqueda y meta del contenido", v(p, "intencion")),
              fila("Palabra clave principal + volumen", "%s (%s)" % (kw, v(p, "volumen"))),
              fila("Palabras clave secundarias + intención", ", ".join(p.get("secundarias") or []) or PEND),
              fila("URL sugerida", v(p, "url"))] + bloque_metas(p) + [fila("H1 sugerido", v(p, "h1", PEND + " (un solo H1 con la keyword)"))]
        L += bloque_nlp(p)
        L += ["", "### II. Imágenes y media", "", "| Elemento | Sugerencias |", "|---|---|", fila("Recomendaciones de imágenes", GUIA_IMG), fila("Otras sugerencias de SEO", GUIA_SEO)]
        L += bloque_links(p) + bloque_estructura(p)
    elif tipo == "landing":
        L += ["### I. Palabras clave", "", "| Elemento | Sugerencias |", "|---|---|", fila("Palabra clave principal + volumen", "%s (%s)" % (kw, v(p, "volumen"))),
              fila("Palabras clave secundarias", ", ".join(p.get("secundarias") or []) or PEND), fila("Top competidores", ", ".join(p.get("competidores") or []) or PEND),
              "", "### II. Información de la página", "", "| Elemento | Sugerencias |", "|---|---|", fila("URL", v(p, "url")), fila("H1", v(p, "h1", PEND + " (un solo H1 con la keyword)"))] + bloque_metas(p)
        L += ["", "### III. Servicio", "", "| Elemento | Sugerencias |", "|---|---|", fila("Servicio principal", v(p, "servicio")), fila("Tipo de servicio", v(p, "tipo_servicio")),
              fila("Propuesta(s) de valor", v(p, "propuesta_valor")), fila("Objetivo de la página", v(p, "objetivo", "Generar leads o contactos (ajusta según el modelo de negocio)")),
              fila("Beneficios (corto, mediano y largo plazo)", v(p, "beneficios")), fila("Prueba social", v(p, "prueba_social")),
              fila("Llamado a la acción", v(p, "cta")), fila("Público objetivo", v(p, "publico")),
              fila("Above the fold", "Propuesta de valor y CTA principal visibles sin hacer scroll")]
        L += bloque_nlp(p) + bloque_links(p) + bloque_estructura(p)
    elif tipo == "ecommerce":
        L += ["### I. Palabras clave (transaccionales)", "", "| Elemento | Sugerencias |", "|---|---|", fila("Palabra clave principal + volumen", "%s (%s)" % (kw, v(p, "volumen"))),
              fila("Palabras clave secundarias", ", ".join(p.get("secundarias") or []) or PEND), fila("Etapa del embudo", v(p, "embudo", "Listo para comprar o comparando opciones")),
              "", "### II. Información de la página", "", "| Elemento | Sugerencias |", "|---|---|", fila("URL", v(p, "url")), fila("H1", v(p, "h1", PEND + " (nombre del producto o categoría con la keyword)"))] + bloque_metas(p)
        L += ["", "### III. Producto o categoría", "", "| Elemento | Sugerencias |", "|---|---|", fila("Atributos técnicos (especificaciones, medidas, materiales)", v(p, "atributos")),
              fila("Beneficios directos (qué problema resuelve)", v(p, "beneficios")), fila("Prueba social (reseñas, valoraciones)", v(p, "prueba_social")),
              fila("Garantías y devoluciones", v(p, "garantias")), fila("Información de envío", v(p, "envio")),
              fila("Urgencia y escasez", v(p, "urgencia")), fila("Cross-sell y up-sell (productos relacionados)", v(p, "cross_sell")),
              fila("Preguntas frecuentes (pagos, envíos, tallas)", v(p, "faq")), fila("Llamado a la acción", v(p, "cta"))]
        L += bloque_nlp(p)
        L += ["", "### IV. Imágenes y media", "", "| Elemento | Sugerencias |", "|---|---|", fila("Recomendaciones de imágenes", GUIA_IMG)]
        L += bloque_links(p) + bloque_estructura(p)
    else:  # actualización (S10): actual vs sugerido
        L += ["### I. Elementos de la página (actual vs sugerido)", "", "| Elemento | Sugerencias |", "|---|---|", fila("Intención de búsqueda", v(p, "intencion")),
              fila("Palabra clave principal + volumen", "%s (%s)" % (kw, v(p, "volumen"))), fila("Posición actual (Search Console)", v(p, "posicion")),
              fila("Palabras clave secundarias", ", ".join(p.get("secundarias") or []) or PEND),
              fila("URL actual", v(p, "url")), fila("URL sugerida", v(p, "url_sugerida", "Mantener la URL actual (si genera tráfico no se cambia; un cambio exige redirección 301)"))] + bloque_metas(p, actual=True)
        L += [fila("H1 actual", v(p, "h1_actual")), fila("H1 sugerido", v(p, "h1", PEND))]
        L += bloque_nlp(p)
        L += ["", "### II. Imágenes y media", "", "| Elemento | Sugerencias |", "|---|---|", fila("Recomendaciones de imágenes", GUIA_IMG), fila("Otras sugerencias de SEO", GUIA_SEO)]
        L += bloque_links(p) + bloque_estructura(p)
        L += ["", "### IV. Cambios a realizar (documentados ANTES de tocar la página)", "", "| Cambio | Prioridad (impacto vs esfuerzo) |", "|---|---|"]
        L += ["| %s | %s |" % (c.get("cambio", c) if isinstance(c, dict) else c, c.get("prioridad", PEND) if isinstance(c, dict) else PEND) for c in p.get("cambios") or []] or ["| %s | %s |" % (PEND, PEND)]
        L += ["", "**Qué se conserva:** %s" % v(p, "conservar", "el núcleo del contenido que ya rankea; los cambios son quirúrgicos (metas, H2, enlaces)"),
              "", "_Prioridad 1: metas y keywords con impresiones. Prioridad 2: schema y reseñas. Prioridad 3: rediseño de imágenes o videos (Semana 10)._"]
    return "\n".join(L) + "\n"


def a_docx(md, ruta):
    from docx import Document  # opcional
    d = Document()
    for l in md.split("\n"):
        if l.startswith("# "):
            d.add_heading(l[2:], 0)
        elif l.startswith("### "):
            d.add_heading(l[4:], 2)
        elif l.startswith("|") and set(l) <= set("|- "):
            continue
        elif l.strip():
            d.add_paragraph(l.strip("| ").replace(" | ", "  ·  "))
    d.save(ruta)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sitio", required=True)
    ap.add_argument("--datos", help="Carpeta de datos (por defecto .seo-audit/<sitio>/data)")
    ap.add_argument("--salida", help="Carpeta de entregables (por defecto .seo-audit/<sitio>/entregables)")
    ap.add_argument("--raiz", default=".")
    ap.add_argument("--docx", action="store_true", help="Además de Markdown, crea .docx (requiere python-docx)")
    args = ap.parse_args()

    base = os.path.join(args.raiz, ".seo-audit", args.sitio)
    datos_dir = args.datos or os.path.join(base, "data")
    salida = args.salida or os.path.join(base, "entregables")
    mapa = leer_json(os.path.join(datos_dir, "content-map.json"))
    if not mapa or not mapa.get("pages"):
        print(json.dumps({"ok": False, "reason": "Falta content-map.json en %s" % datos_dir,
                          "fallback": "corre generar_master.py primero (crea el mapa) o escribe el mapa a mano"}, ensure_ascii=False))
        return 0
    docx_ok = True
    if args.docx:
        try:
            import docx  # noqa: F401
        except ImportError:
            docx_ok = False
    indice, vistos = [], set()
    for p in mapa["pages"]:
        if p.get("accion") not in ("Crear", "Actualizar"):
            continue
        tipo = tipo_brief(p)
        base_n = slug(p.get("keyword") or p.get("url") or "pagina")
        n, i = base_n, 2
        while (tipo, n) in vistos:
            n, i = "%s-%d" % (base_n, i), i + 1
        vistos.add((tipo, n))
        carpeta = asegurar_dir(os.path.join(salida, "briefs", CARPETA[tipo]))
        ruta = os.path.join(carpeta, n + ".md")
        md = brief(p, tipo, args.sitio)
        with open(ruta, "w", encoding="utf-8") as fh:
            fh.write(md)
        if args.docx and docx_ok:
            a_docx(md, ruta[:-3] + ".docx")
        p["brief"] = os.path.relpath(ruta, salida)
        indice.append((p.get("mes", PEND), p.get("keyword", PEND), tipo, p["brief"], p.get("estado", "Pendiente")))
    with open(os.path.join(asegurar_dir(os.path.join(salida, "briefs")), "00-indice.md"), "w", encoding="utf-8") as fh:
        fh.write("# Índice de briefs: %s\n\n| Mes | Keyword | Tipo | Archivo | Estado |\n|---|---|---|---|---|\n" % args.sitio)
        for m, k, t, r, e in indice:
            fh.write("| %s | %s | %s | [%s](%s) | %s |\n" % (m, k, t, os.path.basename(r), r.replace("briefs/", "", 1), e))
    with open(os.path.join(datos_dir, "content-map.json"), "w", encoding="utf-8") as fh:
        json.dump(mapa, fh, ensure_ascii=False, indent=2)
    print(json.dumps({"ok": True, "briefs": len(indice), "salida": os.path.join(salida, "briefs"),
                      "docx": "creado" if args.docx and docx_ok else ("python-docx no instalado: solo Markdown" if args.docx else "no pedido"),
                      "siguiente_paso": "vuelve a correr generar_master.py para que la columna Gdoc apunte a cada brief"}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
