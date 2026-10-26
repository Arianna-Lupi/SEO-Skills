#!/usr/bin/env python3
"""Genera la Plantilla Master (todas sus pestañas) a partir de .seo-audit/<sitio>/data/*.json.

Pestañas y encabezados siguen la Master que comparte la comunidad del diplomado
(ver references/plantillas-comunidad.md): Índice, Herramientas SEO, 1. Plan de SEO,
3. Investigación de Palabras Clave, 5. Mapa de palabras clave, 7. Análisis de Competidores,
9. Estrategia de Contenido, 11. Producción, Actualización de Contenidos,
13. Tareas Revisión Técnica, 15. Auditoría técnica, 17. Plan Especialista SEO, 19. Servicios SEO.

Entrada: carpeta de datos (por defecto .seo-audit/<sitio>/data). Usa si existen:
keywords.json, clusters.json, issues.json, competitors.json, plan-seo.json o next-steps.json
y content-map.json (mapa de contenidos; si falta, se deriva un BORRADOR de keywords+clusters).
Lo que falta queda como `pendiente de dato`: nunca se inventan volumen, KD, clics ni precios.

Salida: <salida>/00-master/Master-<sitio>.xlsx (si hay openpyxl y no usas --csv) o
un CSV por pestaña más LEEME.md para importarlos en Google Sheets.
También escribe content-map.json (con mes y fecha del calendario) y validacion-mapa.json
en la carpeta de datos. Resumen JSON por stdout.

Uso:
  python3 generar_master.py --sitio ejemplo.com --por-mes 4 --inicio 2026-11-02
  python3 generar_master.py --datos ./data --salida ./entregables --csv
Solo stdlib; openpyxl es opcional (pip install openpyxl).
"""
import argparse
import csv
import datetime as dt
import json
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _plan_comun import (ACCIONES, ESTATUS, INTENCIONES, PEND, TIPOS, asegurar_dir, cargar_datos,  # noqa: E402
                         derivar_mapa, ice, pagina_calendario, prioridad, slug, validar_mapa)

SEVERIDAD_A_PRIORIDAD = {"crítico": "Crítico", "critico": "Crítico", "alto": "Alta", "alta": "Alta",
                         "medio": "Media", "media": "Media", "bajo-medio": "Baja", "bajo": "Baja", "baja": "Baja"}

ETIQUETAS_COMPETIDORES = [
    "Sitio web", "Autoridad de Dominio (DR/DA)", "Ranking Ahrefs (AR)", "Dominios referidos (RD)",
    "Tráfico orgánico estimado mensual", "Valor del Tráfico", "¿Tienen un blog?", "Tráfico por País (Ahrefs)",
    "Palabra 1", "Palabra 2", "Palabra 3", "Palabra 4", "Palabra 5",
    "Número de fragmentos destacados (Featured Snippets)",
    "Página Top 1", "Página Top 2", "Página Top 3", "Página Top 4", "Página Top 5",
    "Contenido más enlazado 1", "Contenido más enlazado 2", "Contenido más enlazado 3"]

SERVICIOS = [("SEO en la página", "Análisis y auditoría SEO", "Auditoría Técnica Express", "Una vez"),
             ("SEO en la página", "Investigación de palabras clave", "Keywords, intención y clusters", "Trimestral"),
             ("SEO en la página", "Optimización de contenido", "Pack de Contenidos SEO + IA", "Mensual"),
             ("SEO local", "Gestión de ficha de Google Business Profile", "SEO Local Mensual", "Mensual"),
             ("Reportes", "Reporte mensual con Search Console", "Mantenimiento SEO (retainer)", "Mensual")]


def competidores_top(datos, n=4):
    cuenta = {}
    for g in (datos.get("competitors") or {}).get("by_pillar", []):
        for r in g.get("results", []):
            d = r.get("domain") or ""
            if d:
                cuenta[d] = cuenta.get(d, 0) + 1
    return [d for d, _ in sorted(cuenta.items(), key=lambda x: -x[1])[:n]]


def construir_pestanas(sitio, datos, paginas, avisos):
    T = []  # (nombre, titulo, encabezados, filas, desplegables{col: valores})
    # Índice
    indice = [["Sitio", sitio], ["Generada", dt.date.today().strftime("%d/%m/%Y")],
              ["Estado", "Datos reales cuando existen; el resto dice 'pendiente de dato'"]]
    T.append(("Indice", None, ["Documento", "Detalle"], indice, {}))
    T.append(("Herramientas SEO", "Herramientas Gratuitas y Pagas SEO", ["Area", "Herramienta", "Link", "Uso", "Area"], [
        ["Reportes SEO", "Google Search Console", "https://search.google.com/search-console", "Rendimiento: clics, impresiones, CTR, posición", "Tráfico"],
        ["Reportes SEO", "Google Analytics", "https://analytics.google.com", "Métricas de comportamiento", "Tráfico + usuarios"],
        ["Reportes SEO", "Looker Studio", "https://lookerstudio.google.com", "Automatización de reportes", "Tráfico + usuarios"],
        ["Auditoría", "Screaming Frog (CLI)", "https://www.screamingfrog.co.uk/seo-spider/", "Rastreo e inventario de URLs", "Técnica"],
        ["Auditoría", "PageSpeed Insights", "https://pagespeed.web.dev", "Core Web Vitals", "Velocidad"]], {}))
    # 1. Plan de SEO
    plan = (datos.get("plan-seo") or {}).get("plan")
    filas = []
    if plan:
        for t in plan:
            filas.append(["Mes %s" % t.get("mes"), t.get("tarea"), t.get("detalle") or t.get("observacion") or PEND,
                          t.get("responsable", PEND), t.get("horas", PEND), t.get("fecha", PEND), "Pendiente", t.get("entregable", PEND)])
    else:
        for s in (datos.get("next-steps") or {}).get("steps", []):
            filas.append([PEND, s.get("action"), s.get("impact") or PEND, PEND, s.get("effort", PEND), PEND, "Pendiente", PEND])
    T.append(("1. Plan de SEO", "Plan de SEO", ["Mes", "Tarea", "Objetivo", "Dueño", "Horas", "fecha de Entrega", "Estatus", "Links"],
              filas or [[PEND] * 8], {"Estatus": ESTATUS}))
    # 3. Investigación
    cands = (datos.get("keywords") or {}).get("candidates") or []
    por_kw = {p["keyword"].lower(): p for p in paginas if p.get("keyword")}
    filas = []
    for c in cands:
        p = por_kw.get((c.get("keyword") or "").lower(), {})
        filas.append([c.get("keyword"), (c.get("intent") or "").capitalize() or PEND, p.get("cluster", PEND), p.get("url", PEND),
                      c.get("volume", PEND), c.get("kd", PEND), PEND, p.get("idea", PEND), p.get("h1", PEND), PEND, False])
    T.append(("3. Investigación de Palabras Clave", "Investigación de palabras claves",
              ["Palabra Clave", "Intención de Busqueda", "Cluster", "URL sugerido", "Volúmen de búsqueda", "Dificultad (Keyword Difficulty)",
               "Tasa de Conversion", "Idea de contenido", "H1 sugerido", "Primer Lugar (ejemplo)", "Aprobación del cliente"],
              filas or [[PEND] * 11], {"Intención de Busqueda": INTENCIONES}))
    # 5. Mapa
    filas = []
    for p in paginas:
        filas.append([p.get("url", PEND), p.get("tipo") != "Eliminar", p.get("keyword", PEND), p.get("intencion", PEND), p.get("volumen", PEND),
                      p.get("kd", PEND), p.get("posicion", PEND), p.get("tipo", PEND), {"Crear": "Crear", "Actualizar": "Actualizarla", "Dejar": "Dejarla", "Eliminar": "Eliminarla"}.get(p.get("accion"), PEND),
                      prioridad(p), p.get("notas", ""), "; ".join(p.get("secundarias") or []) or PEND, p.get("cluster", PEND),
                      p.get("impacto", PEND), p.get("confianza", PEND), p.get("facilidad", PEND), ice(p) if ice(p) is not None else PEND,
                      p.get("estado", "Pendiente"), p.get("responsable", PEND), p.get("fecha", PEND)])
    T.append(("5. Mapa de palabras clave", "Modelo de contenido",
              ["url sugerido", "¿Página SEO?", "Palabra clave", "Intención", "Volumen", "Dificultad de la palabra clave", "Posición", "Tipo",
               "¿Dejar, actualizar o eliminar?", "Prioridad", "Notas", "Keywords secundarias", "Cluster", "Impacto", "Confianza", "Facilidad",
               "Puntuación ICE", "Estado", "Responsable", "Fecha calendario"],
              filas or [[PEND] * 20], {"Intención": INTENCIONES, "Tipo": TIPOS, "Estado": ESTATUS}))
    # 7. Competidores
    comp = competidores_top(datos)
    fila_sitio = [ETIQUETAS_COMPETIDORES[0], sitio] + comp
    filas = [fila_sitio] + [[e] + [PEND] * (1 + len(comp)) for e in ETIQUETAS_COMPETIDORES[1:]]
    T.append(("7. Análisis de Competidores", "Reporte de análisis de competidores",
              ["Métrica", "Tu sitio"] + ["Competidor %d" % (i + 1) for i in range(len(comp))], filas, {}))
    # Calendario (9 y 11)
    plan_c = sorted([p for p in paginas if p.get("accion") in ("Crear", "Actualizar")], key=lambda p: (int("".join(c for c in str(p.get("mes", "99")) if c.isdigit()) or 99), -(ice(p) or 0)))
    filas = []
    for i, p in enumerate(plan_c, 1):
        filas.append([i, p.get("mes", PEND), p.get("keyword", PEND), p.get("cluster", PEND), p.get("tipo", PEND), p.get("h1", PEND),
                      p.get("volumen", PEND), PEND, p.get("kd", PEND), PEND, p.get("intencion", PEND), p.get("idea", PEND), PEND, False, p.get("notas", ""),
                      p.get("url", PEND), p.get("estado", "Pendiente")])
    T.append(("9. Estrategia de Contenido", "Estrategia y Plan de contenido",
              ["#", "Mes", "Palabra Clave", "Cluster", "Tipo de contenido", "Titulo H1", "Volumen", "Trafico Potencial", "Dificultad", "Tasa de conversión",
               "Intención de búsqueda", "Idea", "Inspiración #1 en el SERP", "Aprobado?", "Notas", "Pagina", "Estatus"],
              filas or [[PEND] * 17], {"Estatus": ESTATUS, "Tipo de contenido": TIPOS}))
    filas, mes_ant = [], None
    for p in plan_c:
        if p.get("mes") != mes_ant:
            filas.append([str(p.get("mes", PEND)).replace("Mes ", "Mes #")] + [""] * 9)
            mes_ant = p.get("mes")
        filas.append([p.get("keyword", PEND), p.get("tipo", PEND), p.get("cluster", PEND), p.get("url", PEND), p.get("h1", PEND), p.get("estado", "Pendiente"),
                      p.get("fecha", PEND), p.get("brief", PEND), p.get("responsable", PEND), p.get("notas", "")])
    T.append(("11. Produccion", "Producción de contenido",
              ["Palabra Clave", "Tipo de pagina", "Cluster", "URL sugerido", "H1 Sugerido", "Estatus", "Fecha de entrega", "Gdoc", "Autor", "Notas"],
              filas or [[PEND] * 10], {"Estatus": ESTATUS}))
    filas = [[p.get("url"), p.get("keyword"), p.get("posicion", PEND), p.get("meta_title_actual", PEND), p.get("meta_title", PEND),
              p.get("meta_desc_actual", PEND), p.get("meta_desc", PEND), p.get("brief", PEND), p.get("estado", "Pendiente"), ""]
             for p in paginas if p.get("accion") == "Actualizar"]
    T.append(("Actualización de Contenidos", "Actualización de Contenidos o Content Updates",
              ["URL", "Palabra Clave", "Posición", "Meta Titulo Existente", "Meta Titulo Nuevo", "Meta Descripcion Existente", "Meta Descripcion Nueva",
               "Gdoc Contenido nuevo", "Status", "Metas"], filas or [[PEND] * 10], {"Status": ESTATUS}))
    # 13 y 15 desde issues
    issues = (datos.get("issues") or {}).get("issues", [])
    filas = [[i.get("block") or PEND, i.get("title"), (i.get("status") or "Pendiente"), "", (i.get("evidence") or "")[:300], i.get("url", ""), i.get("fix", "")] for i in issues]
    T.append(("13. Tareas Revisión Técnica", None, ["Bloque", "ASPECTO A REVISAR", "ESTATUS", "FOTO", "COMENTARIO", "LINK", "Tasks encontrados"], filas or [[PEND] * 7], {}))
    filas = []
    for i in issues:
        f = {"impacto": i.get("impact"), "confianza": i.get("confidence"), "facilidad": i.get("ease")}
        filas.append([i.get("title"), SEVERIDAD_A_PRIORIDAD.get(str(i.get("severity", "")).lower(), PEND), f["impacto"] or PEND, f["confianza"] or PEND,
                      f["facilidad"] or PEND, ice(f) if ice(f) is not None else PEND, i.get("block") or PEND, i.get("status") or "Pendiente"])
    T.append(("15. Auditoría técnica", "Auditoría técnica SEO",
              ["Recomendación", "Prioridad", "Impacto", "Confianza", "Facilidad", "Puntuación ICE", "Área de impacto", "Estatus"],
              filas or [[PEND] * 8], {"Prioridad": ["Crítico", "Alta", "Media", "Baja"], "Estatus": ESTATUS}))
    areas = ["Estudios y adquisición de conocimientos", "Marca personal", "Perfiles en plataformas de trabajo remoto", "LinkedIn", "Portafolio SEO"]
    T.append(("17. Plan Especialista SEO", "Plan Especialista SEO", ["#", "Área de especialista SEO", "Meta", "Compromiso", "Plazo de tiempo", "Acción"],
              [[i + 1, a, PEND, PEND, PEND, PEND] for i, a in enumerate(areas)], {}))
    T.append(("19. Servicios SEO", sitio, ["Categoría", "Servicio", "Descripción", "Tiempo Estimado", "Frecuencia", "Tiempo promedio", "Precio por Hora", "Precio"],
              [[c, s, d, PEND, f, PEND, PEND, PEND] for c, s, d, f in SERVICIOS], {}))
    return T


def escribir_csv(T, carpeta, sitio):
    asegurar_dir(carpeta)
    nombres = []
    for nombre, titulo, enc, filas, _ in T:
        archivo = "%s.csv" % slug(nombre, 50)
        with open(os.path.join(carpeta, archivo), "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(enc)
            w.writerows(filas)
        nombres.append((nombre, archivo))
    with open(os.path.join(carpeta, "LEEME.md"), "w", encoding="utf-8") as fh:
        fh.write("# Plantilla Master de %s (CSV por pestaña)\n\nNo había openpyxl o pediste --csv. Para llevarla a Google Sheets:\n\n"
                 "1. Crea una copia de la Plantilla Master oficial del diplomado (o una hoja nueva).\n"
                 "2. En cada pestaña usa Archivo, Importar, Subir, y elige 'Reemplazar hoja actual' con el CSV de la lista.\n"
                 "3. Revisa las celdas `pendiente de dato` y complétalas con datos reales (Search Console, Ahrefs, DinoRank).\n\n" % sitio)
        for n, a in nombres:
            fh.write("- %s: `%s`\n" % (n, a))
    return carpeta


def escribir_xlsx(T, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, titulo, enc, filas, desplegables in T:
        ws = wb.create_sheet(nombre[:31])
        r0 = 1
        if titulo:
            ws.cell(1, 1, titulo).font = Font(bold=True, size=13)
            r0 = 2
        for j, h in enumerate(enc, 1):
            c = ws.cell(r0, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, fila in enumerate(filas, r0 + 1):
            for j, v in enumerate(fila, 1):
                ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
        for j in range(1, len(enc) + 1):
            ws.column_dimensions[ws.cell(r0, j).column_letter].width = 28
        ws.freeze_panes = ws.cell(r0 + 1, 1)
        ultima = max(r0 + len(filas), r0 + 200)
        for col, valores in desplegables.items():
            if col in enc:
                dv = DataValidation(type="list", formula1='"%s"' % ",".join(valores), allow_blank=True)
                ws.add_data_validation(dv)
                letra = ws.cell(r0, enc.index(col) + 1).column_letter
                dv.add("%s%d:%s%d" % (letra, r0 + 1, letra, ultima))
    wb.save(ruta)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sitio", required=True, help="Dominio o nombre del proyecto")
    ap.add_argument("--datos", help="Carpeta con los JSON (por defecto .seo-audit/<sitio>/data)")
    ap.add_argument("--salida", help="Carpeta de entregables (por defecto .seo-audit/<sitio>/entregables)")
    ap.add_argument("--raiz", default=".", help="Raíz donde vive .seo-audit")
    ap.add_argument("--por-mes", type=int, default=4, help="Páginas a producir por mes (supuesto editable, por defecto 4)")
    ap.add_argument("--inicio", help="Fecha de inicio AAAA-MM-DD para el calendario (sin ella las fechas quedan pendientes)")
    ap.add_argument("--csv", action="store_true", help="Forzar CSV por pestaña en vez de .xlsx")
    args = ap.parse_args()

    base = os.path.join(args.raiz, ".seo-audit", args.sitio)
    datos_dir = args.datos or os.path.join(base, "data")
    salida = args.salida or os.path.join(base, "entregables")
    datos = cargar_datos(datos_dir)
    if not datos:
        print(json.dumps({"ok": False, "reason": "No hay JSON en %s" % datos_dir,
                          "fallback": "modo manual: copia la Master oficial y llénala con references/plantillas-comunidad.md"}, ensure_ascii=False))
        return 0
    paginas = (datos.get("content-map") or {}).get("pages") or derivar_mapa(datos)
    inicio = dt.datetime.strptime(args.inicio, "%Y-%m-%d").date() if args.inicio else None
    paginas = pagina_calendario(paginas, args.por_mes, inicio)
    avisos = validar_mapa(paginas)

    # Reutiliza el detector existente de canibalización cuando está instalado (no se duplica la lógica).
    canib = None
    script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "mapa-de-palabras-clave", "scripts", "canibalizacion.py")
    if os.path.exists(script):
        tmp = os.path.join(asegurar_dir(datos_dir), "_mapa_tmp.json")
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump([{"url": p.get("url"), "keyword": p.get("keyword"), "volume": p.get("volumen") if isinstance(p.get("volumen"), (int, float)) else None,
                        "position": p.get("posicion") if isinstance(p.get("posicion"), (int, float)) else None} for p in paginas if p.get("keyword")], fh)
        try:
            r = subprocess.run([sys.executable, script, "--file", tmp], capture_output=True, text=True, timeout=30)
            canib = json.loads(r.stdout) if r.stdout.strip() else None
        except (OSError, ValueError, subprocess.SubprocessError):
            canib = None
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)

    T = construir_pestanas(args.sitio, datos, paginas, avisos)
    carpeta_master = asegurar_dir(os.path.join(salida, "00-master"))
    for sub in ("briefs/blog", "briefs/landings", "briefs/ecommerce", "briefs/actualizacion", "reportes"):
        asegurar_dir(os.path.join(salida, sub))
    formato = "csv"
    ruta = carpeta_master
    if not args.csv:
        try:
            ruta = os.path.join(carpeta_master, "Master-%s.xlsx" % slug(args.sitio))
            escribir_xlsx(T, ruta)
            formato = "xlsx"
        except ImportError:
            ruta = carpeta_master
    if formato == "csv":
        escribir_csv(T, carpeta_master, args.sitio)

    with open(os.path.join(datos_dir, "content-map.json"), "w", encoding="utf-8") as fh:
        json.dump({"pages": paginas, "supuestos": {"por_mes": args.por_mes, "inicio": args.inicio or PEND,
                                                    "origen": "content-map.json" if datos.get("content-map") else "borrador derivado de keywords.json y clusters.json"}}, fh, ensure_ascii=False, indent=2)
    with open(os.path.join(datos_dir, "validacion-mapa.json"), "w", encoding="utf-8") as fh:
        json.dump({"avisos": avisos, "canibalizacion": canib}, fh, ensure_ascii=False, indent=2)

    pend = sum(1 for _, _, _, filas, _ in T for f in filas for c in f if c == PEND)
    print(json.dumps({"ok": True, "formato": formato, "ruta": ruta, "pestañas": [t[0] for t in T], "paginas_en_mapa": len(paginas),
                      "origen_mapa": "content-map.json" if datos.get("content-map") else "BORRADOR (revisa y corrige antes de usar)",
                      "avisos_mapa": avisos, "celdas_pendientes": pend, "salida": salida}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
